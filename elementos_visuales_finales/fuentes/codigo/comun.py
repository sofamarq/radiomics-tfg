"""Utilidades comunes para generar las figuras y tablas definitivas de la tesis.

Lee la identidad visual de elementos_visuales_finales/estilo/estilo_tesis.json, registra Source Sans 3
(elementos_visuales_finales/estilo/tipografia/) y define el guardado: figuras en PNG 300 ppp y tablas en CSV
por defecto; PDF/SVG y tablas DOCX o en imagen, opcionales (FIGURAS_FORMATOS, TABLAS_FORMATOS).
Solo lee los resultados del repositorio; no modifica los análisis.
"""
import json
import os
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[3]                 # raíz del repositorio
SALIDA = Path(__file__).resolve().parents[2]               # .../elementos_visuales_finales
DIR_FIG, DIR_TAB, DIR_ANX = SALIDA / "figuras", SALIDA / "tablas", SALIDA / "anexos"
DIR_DATOS = SALIDA / "fuentes" / "datos"
ESTILO = SALIDA / "estilo"
PIPE = RAIZ / "pipeline_v2"
COMPL = RAIZ / "analisis_complementarios" / "resultados"
for d in (DIR_FIG, DIR_TAB, DIR_ANX, DIR_DATOS):
    d.mkdir(parents=True, exist_ok=True)

sys.path.insert(0, str(ESTILO))
import estilo_tesis as et  # noqa: E402

EST = et.EST
TINTA, SUP, COMP, MODELOS, ORDEN = et.TINTA, et.SUP, et.COMP, et.MODELOS, et.ORDEN
CM = et.CM
FAMILIA = "Source Sans 3"
# Formatos de salida: PNG para figuras y CSV para tablas por defecto. Otros formatos ya implementados se
# activan con variables de entorno, p. ej. FIGURAS_FORMATOS="png,pdf,svg" o TABLAS_FORMATOS="csv,docx,png".
FORMATOS_FIG = [f.strip() for f in os.environ.get("FIGURAS_FORMATOS", "png").split(",") if f.strip()]
FORMATOS_TAB = [f.strip() for f in os.environ.get("TABLAS_FORMATOS", "csv").split(",") if f.strip()]

# Colores secundarios (no se confunden con los tres modelos)
SEC = {"pizarra": "#3E4A5C", "violeta": COMP["divergente"]["negativo"], "ocre": "#B07D2B",
       "gris": TINTA["atenuada"]}


def preparar_matplotlib():
    import matplotlib as mpl
    from matplotlib import font_manager
    for ttf in sorted((ESTILO / "tipografia").glob("SourceSans3-*.ttf")):
        font_manager.fontManager.addfont(str(ttf))
    fam = et.aplicar_matplotlib()
    assert fam == FAMILIA, f"Tipografía no disponible: se usaría {fam}"
    mpl.rcParams.update({"svg.fonttype": "path", "pdf.fonttype": 42, "axes.titlesize": 8.5,
                         "legend.handlelength": 2.2, "legend.borderpad": 0.5})
    return mpl


def recortar_png(ruta, margen_px=2):
    """Quita el borde blanco del PNG (deja `margen_px` para no cortar el antialiasing) y conserva los 300 ppp."""
    from PIL import Image, ImageChops
    im = Image.open(ruta)
    rgb = im.convert("RGB")
    caja = ImageChops.difference(rgb, Image.new("RGB", rgb.size, (255, 255, 255))).point(
        lambda v: 255 if v > 8 else 0).getbbox()
    if caja:
        x0, y0, x1, y1 = caja
        caja = (max(x0 - margen_px, 0), max(y0 - margen_px, 0), min(x1 + margen_px, im.width), min(y1 + margen_px, im.height))
        im.crop(caja).save(ruta, dpi=im.info.get("dpi", (300, 300)))


def guardar(fig, nombre):
    """Exporta la figura (por defecto PNG 300 ppp) con el identificador interno como nombre, sin márgenes blancos."""
    for ext in FORMATOS_FIG:
        fig.savefig(DIR_FIG / f"{nombre}.{ext}", **{**et.GUARDADO, "pad_inches": 0.01})
    if "png" in FORMATOS_FIG:
        recortar_png(DIR_FIG / f"{nombre}.png")


def letra_panel(ax, letra, x=-0.16, y=1.04):
    ax.text(x, y, letra, transform=ax.transAxes, fontsize=EST["tipografia"]["tamanos_pt_matplotlib"]["letra_panel"] + 1,
            fontweight="bold", color=TINTA["primaria"], va="bottom", ha="left")


num = et.num


def pct(x, d=1):
    return f"{x * 100:.{d}f}".replace(".", ",") + " %"


def ic(lo, hi, d=3):
    return f"[{num(lo, d)}; {num(hi, d)}]"


# ---------------------------------------------------------------- tablas DOCX
def _borde(celda, lado, grosor_pt, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = celda._tc.get_or_add_tcPr()
    bordes = tcPr.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        tcPr.append(bordes)
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(int(round(grosor_pt * 8))))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color.lstrip("#"))
    bordes.append(el)


def _preparar_tabla(tabla, anchos_cm):
    """Ancho fijo por columna (tblW, tblGrid, tblLayout fixed) y bordes de tabla nulos, respetando el orden
    de elementos de tblPr exigido por el esquema de Word (tblW, jc, tblBorders, tblLayout, tblCellMar, tblLook)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = tabla._tbl.tblPr
    total = int(round(sum(anchos_cm) / 2.54 * 1440))
    tblW = tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(total))
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("left", 85), ("right", 85)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    look = tblPr.find(qn("w:tblLook"))
    for el in (b, lay, mar):
        look.addprevious(el)
    for gc, w in zip(tabla._tbl.tblGrid.findall(qn("w:gridCol")), anchos_cm):
        gc.set(qn("w:w"), str(int(round(w / 2.54 * 1440))))


def _estilo_base(doc, tam):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    st = doc.styles["Normal"]
    st.font.name = FAMILIA
    st.font.size = Pt(tam)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FAMILIA)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0


def _fuente(run, tam, negrita=False, cursiva=False, color=None):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    run.font.name = FAMILIA
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FAMILIA)
    run.font.size = Pt(tam)
    run.font.bold = negrita
    run.font.italic = cursiva
    c = (color or TINTA["primaria"]).lstrip("#")
    run.font.color.rgb = RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))


def _no_separar(txt):
    """Evita cortes de línea poco legibles: «ROC-AUC», «67,7 %», «n = 93», «C = 0,1»."""
    import re
    txt = txt.replace("ROC-AUC", "ROC\u2011AUC").replace(" %", "\u00a0%")
    return re.sub(r"(\S) = (\S)", "\\1\u00a0=\u00a0\\2", txt)


def tabla_docx(df, ruta, alinear_der=None, nota=None, anchos_cm=None, grupos=None, negrita_filas=None,
               tam=9.0, encabezado_superior=None):
    """Tabla de tres filetes (booktabs): filete superior e inferior gruesos, filete bajo el encabezado,
    separadores de fila finos y grises, sin bordes verticales ni sombreado. Sin número ni epígrafe
    (se agregan en el documento). `grupos`: índices de filas que inician un grupo (etiqueta en negrita
    en la primera columna). `negrita_filas`: filas en negrita. `encabezado_superior`: lista de
    (texto, columnas que abarca) para un primer nivel de encabezado con filete parcial debajo."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt

    t = EST["tablas"]
    gris = EST["tablas"]["separador_filas"]["color"]
    doc = Document()
    _estilo_base(doc, tam)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    alinear_der = set(alinear_der or [])
    negrita_filas = set(negrita_filas or [])
    grupos = set(grupos or [])
    nfil, ncol = df.shape
    h = 2 if encabezado_superior else 1
    tabla = doc.add_table(rows=nfil + h, cols=ncol)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    anchos_cm = anchos_cm or [16.0 / ncol] * ncol
    assert sum(anchos_cm) <= 16.05, f"tabla más ancha que la caja de texto: {sum(anchos_cm):.2f} cm"
    _preparar_tabla(tabla, anchos_cm)
    al = lambda j: WD_ALIGN_PARAGRAPH.RIGHT if j in alinear_der else WD_ALIGN_PARAGRAPH.LEFT  # noqa: E731

    def escribir(c, txt, negrita, alin, sangria=False):
        c.text = ""
        p = c.paragraphs[0]
        if sangria:
            p.paragraph_format.left_indent = Cm(0.35)
        _fuente(p.add_run(_no_separar(txt)), tam, negrita=negrita)
        p.alignment = alin

    if encabezado_superior:
        assert sum(n for _, n in encabezado_superior) == ncol
        j = 0
        for texto, n in encabezado_superior:
            c = tabla.cell(0, j)
            if n > 1:
                c = c.merge(tabla.cell(0, j + n - 1))
            escribir(c, texto, True, WD_ALIGN_PARAGRAPH.CENTER if n > 1 else al(j))
            c.width = Cm(sum(anchos_cm[j:j + n]))
            _borde(c, "top", t["filete_superior_inferior"], TINTA["primaria"])
            if texto and n > 1:
                _borde(c, "bottom", t["separador_filas"]["ancho"] * 1.4, TINTA["eje"])
            j += n
    for j, col in enumerate(df.columns):
        c = tabla.cell(h - 1, j)
        escribir(c, str(col), True, al(j))
        if h == 1:
            _borde(c, "top", t["filete_superior_inferior"], TINTA["primaria"])
        _borde(c, "bottom", t["filete_encabezado"], TINTA["primaria"])
    for i in range(nfil):
        for j in range(ncol):
            v = df.iat[i, j]
            c = tabla.cell(i + h, j)
            txt = "" if v is None or (isinstance(v, float) and v != v) else str(v)
            sangria = j == 0 and txt.startswith("   ")        # categoría dentro de un grupo
            escribir(c, txt.strip() if sangria else txt, (i in negrita_filas) or (i in grupos and j == 0), al(j),
                     sangria)
            if i == nfil - 1:
                _borde(c, "bottom", t["filete_superior_inferior"], TINTA["primaria"])
            elif (i + 1) in grupos:
                _borde(c, "bottom", t["filete_encabezado"] * 0.75, TINTA["eje"])
            else:
                _borde(c, "bottom", t["separador_filas"]["ancho"], gris)
    for fila in tabla.rows:
        for c in fila.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
    for j, w in enumerate(anchos_cm):
        for fila in tabla.rows[h - 1:]:
            fila.cells[j].width = Cm(w)
    for fila in tabla.rows[:h]:                       # encabezado repetido si la tabla cambia de página
        fila._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for fila in tabla.rows:                           # las filas no se parten entre páginas
        fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    if nota:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        _fuente(p.add_run("Nota. "), tam - 1, cursiva=True, color=TINTA["secundaria"])
        _fuente(p.add_run(_no_separar(nota)), tam - 1, color=TINTA["secundaria"])
    doc.save(ruta)


def tabla_imagen(df, ruta_base, alinear_der=None, nota=None, anchos_cm=None, grupos=None, negrita_filas=None,
                 tam=9.0, encabezado_superior=None):
    """Misma tabla de tres filetes que tabla_docx, dibujada con matplotlib (PNG 300 ppp, PDF y SVG).
    Las medidas se trabajan en puntos tipográficos; el texto se ajusta al ancho de cada columna."""
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    from matplotlib.textpath import TextToPath
    preparar_matplotlib()
    t = EST["tablas"]
    ttp = TextToPath()
    fp = {b: FontProperties(family=FAMILIA, weight="bold" if b else "normal", size=tam) for b in (False, True)}
    ancho = lambda txt, b=False, prop=None: ttp.get_text_width_height_descent(txt, prop or fp[b], ismath=False)[0]  # noqa: E731

    def ajustar(txt, disp, b=False, prop=None):
        lineas, act = [], ""
        for parrafo in str(txt).split("\n"):
            for pal in parrafo.split(" "):
                prueba = f"{act} {pal}" if act else pal
                if act and ancho(prueba, b, prop) > disp:
                    lineas.append(act)
                    act = pal
                else:
                    act = prueba
            lineas.append(act)
            act = ""
        return lineas

    PT = 72 / 2.54
    alinear_der, grupos, negrita_filas = set(alinear_der or []), set(grupos or []), set(negrita_filas or [])
    nfil, ncol = df.shape
    anchos = [w * PT for w in (anchos_cm or [16.0 / ncol] * ncol)]
    x0 = [sum(anchos[:j]) for j in range(ncol)]
    W, pad, sang, inter, vpad = sum(anchos), 0.10 * PT, 0.35 * PT, tam * 1.22, 3.2
    elems, y = [], 0.0          # y crece hacia abajo, en puntos
    reglas = []                 # (y, x_ini, x_fin, grosor, color)
    fil_sup, fil_enc = t["filete_superior_inferior"] * 0.7, t["filete_encabezado"] * 0.7
    sep, gris = t["separador_filas"]["ancho"] * 0.7, t["separador_filas"]["color"]
    reglas.append((0.0, 0, W, fil_sup, TINTA["primaria"]))

    def fila(celdas, negritas, alins, sangrias, y):
        bloques = []
        for j, txt in enumerate(celdas):
            disp = anchos[j] - 2 * pad - (sang if sangrias[j] else 0)
            bloques.append(ajustar(txt, disp, negritas[j]) if txt else [""])
        alto = max(len(b) for b in bloques) * inter + 2 * vpad
        for j, lineas in enumerate(bloques):
            for k, ln in enumerate(lineas):
                xi = x0[j] + pad + (sang if sangrias[j] else 0)
                xx = x0[j] + anchos[j] - pad if alins[j] == "right" else (x0[j] + anchos[j] / 2 if alins[j] == "center" else xi)
                elems.append((xx, y + vpad + (k + 0.78) * inter, ln, negritas[j], alins[j]))
        return y + alto

    al = lambda j: "right" if j in alinear_der else "left"  # noqa: E731
    if encabezado_superior:
        ini, bloques = 0, []
        for texto, n in encabezado_superior:
            xa, xb = x0[ini], x0[ini] + sum(anchos[ini:ini + n])
            bloques.append((texto, n, ini, xa, xb, ajustar(texto, xb - xa - 2 * pad, True) if texto else [""]))
            ini += n
        y1 = max(len(b[5]) for b in bloques) * inter + 2 * vpad
        for texto, n, ini, xa, xb, lineas in bloques:
            a = "center" if n > 1 else al(ini)
            xx = (xa + xb) / 2 if n > 1 else (xb - pad if a == "right" else xa + pad)
            for m, ln in enumerate(lineas):
                elems.append((xx, vpad + (m + 0.78) * inter, ln, True, a))
            if texto and n > 1:
                reglas.append((y1, xa + 0.15 * PT, xb - 0.15 * PT, sep * 1.4, TINTA["eje"]))
        y = y1
    y = fila([str(c).strip() for c in df.columns], [True] * ncol, [al(j) for j in range(ncol)], [False] * ncol, y)
    reglas.append((y, 0, W, fil_enc, TINTA["primaria"]))
    for i in range(nfil):
        celdas, sangrias = [], []
        for j in range(ncol):
            v = df.iat[i, j]
            txt = "" if v is None or (isinstance(v, float) and v != v) else str(v)
            sangrias.append(j == 0 and txt.startswith("   "))
            celdas.append(txt.strip() if sangrias[-1] else txt)
        negr = [(i in negrita_filas) or (i in grupos and j == 0) for j in range(ncol)]
        y = fila(celdas, negr, [al(j) for j in range(ncol)], sangrias, y)
        if i == nfil - 1:
            reglas.append((y, 0, W, fil_sup, TINTA["primaria"]))
        elif (i + 1) in grupos:
            reglas.append((y, 0, W, fil_enc * 0.75, TINTA["eje"]))
        else:
            reglas.append((y, 0, W, sep, gris))
    lineas_nota = []
    if nota:
        fpn = {c: FontProperties(family=FAMILIA, style="italic" if c else "normal", size=tam - 1) for c in (False, True)}
        pref = "Nota. "
        lineas_nota = ajustar(pref + nota, W, prop=fpn[False])
        y_nota = y + 5.0
        y += 5.0 + len(lineas_nota) * (tam - 1) * 1.25
    H = y + 1.0
    fig = plt.figure(figsize=(W / 72, H / 72))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(H, 0)
    ax.axis("off")
    for yy, xa, xb, lw, col in reglas:
        ax.plot([xa, xb], [yy, yy], color=col, linewidth=lw, solid_capstyle="butt")
    for xx, yy, ln, b, a in elems:
        ax.text(xx, yy, ln, ha=a, va="baseline", fontsize=tam, fontweight="bold" if b else "normal",
                color=TINTA["primaria"], fontfamily=FAMILIA)
    for k, ln in enumerate(lineas_nota):
        yy = y_nota + (k + 0.8) * (tam - 1) * 1.25
        if k == 0 and ln.startswith(pref):
            ax.text(0, yy, "Nota.", fontsize=tam - 1, style="italic", color=TINTA["secundaria"], va="baseline")
            ax.text(ancho(pref, prop=fpn[True]), yy, ln[len(pref):], fontsize=tam - 1, color=TINTA["secundaria"],
                    va="baseline")
        else:
            ax.text(0, yy, ln, fontsize=tam - 1, color=TINTA["secundaria"], va="baseline")
    for ext in ("png", "pdf", "svg"):
        fig.savefig(f"{ruta_base}.{ext}", dpi=EST["dimensiones_A4_cm"]["dpi_exportacion"], facecolor=SUP["fondo"],
                    bbox_inches="tight", pad_inches=0.01)
    plt.close(fig)
    recortar_png(f"{ruta_base}.png")


def guardar_tabla(df, ident, destino="tablas", **kw):
    """Guarda la tabla en CSV (UTF-8); opcionalmente, en .docx y como imagen (.png, .pdf y .svg) con el mismo
    diseño, según TABLAS_FORMATOS."""
    base = {"tablas": DIR_TAB, "anexos": DIR_ANX}[destino]
    df.to_csv(base / f"{ident}.csv", index=False, encoding="utf-8")
    if "docx" in FORMATOS_TAB:
        tabla_docx(df, base / f"{ident}.docx", **kw)
    if {"png", "pdf", "svg"} & set(FORMATOS_TAB):
        tabla_imagen(df, base / ident, **kw)


def leer_json(ruta):
    return json.loads(Path(ruta).read_text(encoding="utf-8"))
